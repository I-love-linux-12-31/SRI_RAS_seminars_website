"""Проверки панели управления."""

from datetime import datetime, time, timedelta

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.urls import reverse
from django.utils import timezone

from apps.registrations.models import Registration
from apps.seminars.models import Seminar, Speaker
from apps.seminars.tests.factories import add_talk, make_seminar

pytestmark = pytest.mark.django_db


@pytest.fixture
def groups():
    call_command("setup_groups")


@pytest.fixture
def secretary(django_user_model, groups):
    from django.contrib.auth.models import Group

    user = django_user_model.objects.create_user(username="sec", password="pw", is_staff=True)
    user.groups.add(Group.objects.get(name="Секретарь семинара"))
    return user


@pytest.fixture
def chair(django_user_model, groups):
    from django.contrib.auth.models import Group

    user = django_user_model.objects.create_user(username="chair", password="pw", is_staff=True)
    user.groups.add(Group.objects.get(name="Руководитель семинара"))
    return user


@pytest.fixture
def as_secretary(client, secretary):
    client.force_login(secretary)
    return client


def make_registration(seminar, **kwargs) -> Registration:
    data = {
        "seminar": seminar,
        "full_name": "Иванов Иван Иванович",
        "organization": "МГУ",
        "position": "аспирант",
        "email": "ivanov@msu.ru",
        "consent_pd": True,
        "consent_at": timezone.now(),
    }
    return Registration.objects.create(**(data | kwargs))


# --- Доступ -------------------------------------------------------------------


PANEL_URLS = [
    ("staffpanel:seminar_list", {}),
    ("staffpanel:registrations", {}),
    ("staffpanel:settings", {}),
]


@pytest.mark.parametrize(("name", "kwargs"), PANEL_URLS)
def test_anonymous_is_sent_to_login(client, name, kwargs):
    response = client.get(reverse(name, kwargs=kwargs))

    assert response.status_code == 302
    assert reverse("staffpanel:login") in response["Location"]


@pytest.mark.parametrize(("name", "kwargs"), PANEL_URLS)
def test_plain_user_gets_403(client, django_user_model, name, kwargs):
    django_user_model.objects.create_user(username="guest", password="pw")
    client.force_login(django_user_model.objects.get(username="guest"))

    assert client.get(reverse(name, kwargs=kwargs)).status_code == 403


def test_secretary_cannot_edit_site_settings(as_secretary):
    """Формулировку согласия на ПД меняет руководитель, а не любой сотрудник."""
    assert as_secretary.get(reverse("staffpanel:settings")).status_code == 403


def test_chair_can_edit_site_settings(client, chair):
    client.force_login(chair)

    assert client.get(reverse("staffpanel:settings")).status_code == 200


def test_secretary_can_manage_seminars(as_secretary):
    assert as_secretary.get(reverse("staffpanel:seminar_create")).status_code == 200
    assert as_secretary.get(reverse("staffpanel:registrations")).status_code == 200


# --- Список -------------------------------------------------------------------


def test_list_shows_all_statuses_including_drafts(as_secretary):
    draft = make_seminar(timezone.localdate(), suffix="d", status=Seminar.Status.DRAFT)
    add_talk(draft, "Доклад в черновике", [("Иванов И. И.", "ИКИ РАН")])
    hidden = make_seminar(timezone.localdate(), suffix="h", status=Seminar.Status.HIDDEN)
    add_talk(hidden, "Доклад в скрытом", [("Петров П. П.", "ИКИ РАН")])

    content = as_secretary.get(reverse("staffpanel:seminar_list")).content.decode()

    assert "Доклад в черновике" in content
    assert "Доклад в скрытом" in content


def test_list_identifies_seminars_by_talk_and_speaker(as_secretary):
    """Тематики больше нет: в управлении заседание опознаётся по докладу."""
    seminar = make_seminar(timezone.localdate() + timedelta(days=5))
    add_talk(seminar, "Турбулентность солнечного ветра", [("Смирнов Андрей", "ИКИ РАН")])

    content = as_secretary.get(reverse("staffpanel:seminar_list")).content.decode()

    assert "Турбулентность солнечного ветра" in content
    assert "Смирнов Андрей" in content


def test_list_counts_registrations(as_secretary):
    seminar = make_seminar(timezone.localdate() + timedelta(days=5))
    make_registration(seminar, email="a@example.org")
    make_registration(seminar, email="b@example.org")

    content = as_secretary.get(reverse("staffpanel:seminar_list")).content.decode()

    assert f'?seminar={seminar.pk}">2</a>' in content


# --- Создание и правка --------------------------------------------------------


def seminar_payload(**overrides) -> dict:
    data = {
        "date": (timezone.localdate() + timedelta(days=20)).isoformat(),
        "start_time": "11:00",
        "abstract_ru": "",
        "abstract_en": "",
        "place_ru": "ИКИ РАН",
        "place_en": "",
        "seminar_format": Seminar.Format.HYBRID,
        "online_url": "",
        "status": Seminar.Status.DRAFT,
        "registration_closes_at": "",
        "talks-TOTAL_FORMS": "1",
        "talks-INITIAL_FORMS": "0",
        "talks-MIN_NUM_FORMS": "0",
        "talks-MAX_NUM_FORMS": "1000",
        "talks-0-title_ru": "Турбулентность солнечного ветра",
        "talks-0-title_en": "",
        "talks-0-abstract_ru": "",
        "talks-0-abstract_en": "",
        "talks-0-speakers_raw": "Смирнов Андрей Петрович — д. ф.-м. н., ИКИ РАН",
    }
    return data | overrides


def saved_talk_payload(talk, **overrides) -> dict:
    """Правка заседания с одним уже сохранённым докладом."""
    return seminar_payload(
        **{
            "talks-INITIAL_FORMS": "1",
            "talks-0-id": talk.pk,
            "talks-0-title_ru": talk.title_ru,
            "talks-0-speakers_raw": "Иванов Иван Иванович — ИКИ РАН",
        }
        | overrides
    )


def material_payload(talk, **fields) -> dict:
    """Формсет материалов идёт своей группой на каждый доклад.

    Без полей отдаёт один management_form: у формы материала есть значение
    по умолчанию в поле «тип», и пустая лишняя строка считалась бы заполненной.
    """
    prefix = f"materials-{talk.pk}"
    data = {
        f"{prefix}-TOTAL_FORMS": "1" if fields else "0",
        f"{prefix}-INITIAL_FORMS": "0",
        f"{prefix}-MIN_NUM_FORMS": "0",
        f"{prefix}-MAX_NUM_FORMS": "1000",
    }
    return data | {f"{prefix}-0-{name}": value for name, value in fields.items()}


def test_create_seminar_with_talk_and_speaker(as_secretary):
    response = as_secretary.post(reverse("staffpanel:seminar_create"), seminar_payload())

    assert response.status_code == 302
    seminar = Seminar.objects.get()
    assert seminar.slug, "slug должен проставляться автоматически"
    talk = seminar.talks.get()
    assert [s.full_name_ru for s in talk.ordered_speakers] == ["Смирнов Андрей Петрович"]
    assert Speaker.objects.get().affiliation_ru == "д. ф.-м. н., ИКИ РАН"


def test_talk_abstract_is_saved_in_both_languages(as_secretary):
    """Английская аннотация доклада необязательна, но панель должна её принимать."""
    payload = seminar_payload(
        **{
            "talks-0-abstract_ru": "Аннотация доклада по-русски.",
            "talks-0-abstract_en": "The talk abstract in English.",
        }
    )

    as_secretary.post(reverse("staffpanel:seminar_create"), payload)

    talk = Seminar.objects.get().talks.get()
    assert talk.abstract_ru == "Аннотация доклада по-русски."
    assert talk.abstract_en == "The talk abstract in English."


def test_talk_saves_without_the_english_abstract(as_secretary):
    response = as_secretary.post(reverse("staffpanel:seminar_create"), seminar_payload())

    assert response.status_code == 302
    assert Seminar.objects.get().talks.get().abstract_en == ""


def test_speakers_are_reused_not_duplicated(as_secretary):
    """Один и тот же докладчик выступает годами — справочник плодиться не должен."""
    as_secretary.post(reverse("staffpanel:seminar_create"), seminar_payload())
    as_secretary.post(
        reverse("staffpanel:seminar_create"),
        seminar_payload(**{"talks-0-title_ru": "Второй доклад"}),
    )

    assert Speaker.objects.count() == 1
    assert Seminar.objects.count() == 2


def test_speaker_order_is_kept_from_input(as_secretary):
    payload = seminar_payload(
        **{"talks-0-speakers_raw": "Чернецов Никита Севирович\nКавокин Кирилл Витальевич"},
    )

    as_secretary.post(reverse("staffpanel:seminar_create"), payload)

    talk = Seminar.objects.get().talks.get()
    assert [s.full_name_ru for s in talk.ordered_speakers] == [
        "Чернецов Никита Севирович",
        "Кавокин Кирилл Витальевич",
    ]


def test_new_seminar_feeds_search_index(as_secretary):
    """Доклад сохраняется после заседания — поисковый текст надо пересобрать."""
    as_secretary.post(reverse("staffpanel:seminar_create"), seminar_payload())

    seminar = Seminar.objects.get()
    assert "Турбулентность солнечного ветра" in seminar.search_text
    assert "Смирнов Андрей Петрович" in seminar.search_text


def test_online_seminar_rejects_pass_requirement(as_secretary):
    payload = seminar_payload(seminar_format=Seminar.Format.ONLINE, pass_required="on")

    response = as_secretary.post(reverse("staffpanel:seminar_create"), payload)

    assert response.status_code == 422
    assert Seminar.objects.count() == 0
    assert "пропуск не нужен" in response.content.decode()


def test_invalid_form_does_not_create_anything(as_secretary):
    response = as_secretary.post(
        reverse("staffpanel:seminar_create"), seminar_payload(**{"talks-0-title_ru": ""})
    )

    assert response.status_code == 422
    assert Seminar.objects.count() == 0
    assert Speaker.objects.count() == 0, "докладчики не должны создаваться при откате"


def test_edit_form_prefills_date_and_time(as_secretary):
    """Русская локаль печатает «17.08.2026», а <input type="date"> ждёт ISO.

    Без явного формата поле при правке открывалось пустым, и заседание
    сохранялось без даты — точнее, не сохранялось вовсе.
    """
    seminar = make_seminar(
        timezone.localdate() + timedelta(days=20),
        start_time=time(15, 30),
        registration_closes_at=timezone.make_aware(
            datetime.combine(timezone.localdate() + timedelta(days=18), time(18, 0))
        ),
    )

    content = as_secretary.get(
        reverse("staffpanel:seminar_edit", kwargs={"pk": seminar.pk})
    ).content.decode()

    assert f'value="{seminar.date:%Y-%m-%d}"' in content
    assert 'value="15:30"' in content
    assert f'value="{seminar.date - timedelta(days=2):%Y-%m-%d}T18:00"' in content


# --- Сводка ошибок --------------------------------------------------------------


def test_invalid_form_lists_problems_above_the_form(as_secretary):
    response = as_secretary.post(
        reverse("staffpanel:seminar_create"),
        seminar_payload(date="", **{"talks-0-title_ru": ""}),
    )

    content = response.content.decode()
    assert "Не сохранено" in content
    assert "Дата: Обязательное поле." in response.context["errors"]
    assert "Название доклада: Обязательное поле." in response.context["errors"]


def test_summary_collects_errors_from_talks_and_materials(as_secretary):
    """Ошибка во вложенной форме иначе видна только при прокрутке до неё."""
    seminar = make_seminar(timezone.localdate() + timedelta(days=5))
    talk = add_talk(seminar, "Доклад", [("Иванов Иван Иванович", "ИКИ РАН")])
    payload = saved_talk_payload(talk, **{"talks-0-speakers_raw": "АБ"})
    payload |= material_payload(talk, kind="slides", title_ru="Презентация", url="")

    response = as_secretary.post(
        reverse("staffpanel:seminar_edit", kwargs={"pk": seminar.pk}), payload
    )

    errors = response.context["errors"]
    assert "Докладчики: Слишком короткое имя докладчика: «АБ»" in errors
    assert "Приложите файл или укажите ссылку." in errors


# --- Материалы доклада ----------------------------------------------------------


def talk_with_seminar():
    seminar = make_seminar(timezone.localdate() + timedelta(days=5))
    return seminar, add_talk(seminar, "Доклад", [("Иванов Иван Иванович", "ИКИ РАН")])


def test_talk_is_deleted_through_the_formset_checkbox(as_secretary):
    """Кнопка удаления — это обёртка над чекбоксом формсета, и он должен работать."""
    seminar = make_seminar(timezone.localdate() + timedelta(days=20))
    talk = add_talk(seminar, "Лишний доклад", [("Иванов И. И.", "ИКИ РАН")])
    keeper = add_talk(seminar, "Нужный доклад", [("Петров П. П.", "ИКИ РАН")])

    payload = seminar_payload(
        **{
            "talks-TOTAL_FORMS": "2",
            "talks-INITIAL_FORMS": "2",
            "talks-0-id": talk.pk,
            "talks-0-title_ru": talk.title_ru,
            "talks-0-speakers_raw": "Иванов Иван Иванович — ИКИ РАН",
            "talks-0-DELETE": "on",
            "talks-1-id": keeper.pk,
            "talks-1-title_ru": keeper.title_ru,
            "talks-1-title_en": "",
            "talks-1-abstract_ru": "",
            "talks-1-abstract_en": "",
            "talks-1-speakers_raw": "Петров Пётр Петрович — ИКИ РАН",
        }
    )

    # У каждого сохранённого доклада своя группа материалов — без её
    # management_form форма не проходит проверку.
    payload |= material_payload(talk) | material_payload(keeper)

    response = as_secretary.post(
        reverse("staffpanel:seminar_edit", kwargs={"pk": seminar.pk}), payload
    )

    assert response.status_code == 302
    assert [t.title_ru for t in seminar.talks.all()] == ["Нужный доклад"]


def test_delete_button_keeps_the_checkbox_and_asks_before_deleting(as_secretary):
    """Заказчик просил кнопку с подтверждением вместо галочки.

    Чекбокс формсета остаётся в разметке — без него Django об удалении не
    узнает и панель перестанет работать без JavaScript, — но рядом стоит
    кнопка, а вопрос для неё готовит шаблон.
    """
    seminar = make_seminar(timezone.localdate() + timedelta(days=20))
    add_talk(seminar, "Турбулентность солнечного ветра", [("Иванов И. И.", "ИКИ РАН")])

    content = as_secretary.get(
        reverse("staffpanel:seminar_edit", kwargs={"pk": seminar.pk})
    ).content.decode()

    assert 'name="talks-0-DELETE"' in content
    assert "data-delete-inline" in content
    assert "Удалить доклад «Турбулентность солнечного ветра»?" in content


def test_material_is_attached_to_the_talk(as_secretary):
    """Заказчик просил файлы у доклада, а не у заседания."""
    from apps.seminars.models import Material

    seminar, talk = talk_with_seminar()
    payload = saved_talk_payload(talk)
    payload |= material_payload(
        talk,
        kind=Material.Kind.SLIDES,
        title_ru="Презентация",
        file=SimpleUploadedFile("slides.pdf", b"%PDF-1.4"),
        url="",
    )

    response = as_secretary.post(
        reverse("staffpanel:seminar_edit", kwargs={"pk": seminar.pk}), payload
    )

    assert response.status_code == 302
    material = Material.objects.get()
    assert material.talk_id == talk.pk
    assert material.file.name.startswith("materials/")


def test_material_can_be_a_link(as_secretary):
    from apps.seminars.models import Material

    seminar, talk = talk_with_seminar()
    payload = saved_talk_payload(talk)
    payload |= material_payload(
        talk, kind=Material.Kind.VIDEO, title_ru="", url="https://example.org/video"
    )

    response = as_secretary.post(
        reverse("staffpanel:seminar_edit", kwargs={"pk": seminar.pk}), payload
    )

    assert response.status_code == 302
    assert Material.objects.get().url == "https://example.org/video"


def test_new_seminar_has_nowhere_to_attach_materials_yet(as_secretary):
    """Материал висит на докладе, а доклада у несохранённой формы ещё нет."""
    content = as_secretary.get(reverse("staffpanel:seminar_create")).content.decode()

    assert "materials-" not in content
    assert "после сохранения" in content


def test_material_form_appears_for_a_saved_talk(as_secretary):
    seminar, talk = talk_with_seminar()

    content = as_secretary.get(
        reverse("staffpanel:seminar_edit", kwargs={"pk": seminar.pk})
    ).content.decode()

    assert f'name="materials-{talk.pk}-0-file"' in content
    assert f'name="materials-{talk.pk}-0-url"' in content


def test_saved_form_has_no_summary(as_secretary):
    content = as_secretary.get(reverse("staffpanel:seminar_create")).content.decode()

    assert "formsummary" not in content


# --- Настройки сайта ------------------------------------------------------------


def settings_payload(**overrides) -> dict:
    """Заполненная форма настроек: пустыми обязательные поля оставлять нельзя."""
    data = {
        "chair_ru": "Зеленый Лев Матвеевич",
        "chair_en": "",
        "chair_title_ru": "академик РАН",
        "chair_title_en": "",
        "secretary_ru": "Евдокимова Дарья Геннадьевна",
        "secretary_en": "",
        "secretary_title_ru": "к. ф.-м. н.",
        "secretary_title_en": "",
        "email": "seminar@cosmos.ru",
        "address_ru": "Москва, ул. Профсоюзная, 84/32",
        "address_en": "",
        "registration_lead_hours": "48",
        "foreign_extra_lead_hours": "48",
        "notify_on_registration": "on",
        "notify_email": "",
        "online_link_captcha": "on",
        "retention_months": "6",
    }
    return data | overrides


def pdf_upload(name: str = "soglasie.pdf") -> SimpleUploadedFile:
    return SimpleUploadedFile(name, b"%PDF-1.4\ntrailer\n%%EOF\n", content_type="application/pdf")


def test_chair_uploads_the_consent_document(client, chair):
    """Согласие на ПД — единственное, что приходит документом, а не текстом."""
    from apps.core.models import SiteSettings

    client.force_login(chair)

    response = client.post(
        reverse("staffpanel:settings"), settings_payload(privacy_policy_file=pdf_upload())
    )

    assert response.status_code == 302
    assert SiteSettings.objects.get(pk=1).privacy_policy_file.name.startswith("policy/")


def test_consent_document_must_be_a_pdf(client, chair):
    """Документ открывается по ссылке прямо в браузере — значит, PDF."""
    from apps.core.models import SiteSettings

    client.force_login(chair)

    response = client.post(
        reverse("staffpanel:settings"),
        settings_payload(privacy_policy_file=SimpleUploadedFile("soglasie.docx", b"PK\x03\x04")),
    )

    assert response.status_code == 422
    assert not SiteSettings.objects.get(pk=1).privacy_policy_file


# --- Клонирование -------------------------------------------------------------


def test_clone_copies_talks_but_not_registrations(as_secretary):
    source = make_seminar(
        timezone.localdate() - timedelta(days=30), status=Seminar.Status.PUBLISHED
    )
    add_talk(source, "Прошлый доклад", [("Иванов И. И.", "ИКИ РАН")])
    make_registration(source)

    response = as_secretary.post(reverse("staffpanel:seminar_clone", kwargs={"pk": source.pk}))

    assert response.status_code == 302
    clone = Seminar.objects.exclude(pk=source.pk).get()
    assert clone.status == Seminar.Status.DRAFT, "копия не должна сразу уходить в публикацию"
    assert clone.talks.count() == 1
    assert clone.registrations.count() == 0, "чужие заявки копировать нельзя"
    assert Registration.objects.count() == 1


def test_clone_keeps_original_intact(as_secretary):
    source = make_seminar(timezone.localdate() - timedelta(days=30))
    add_talk(source, "Прошлый доклад", [("Иванов И. И.", "ИКИ РАН")])

    as_secretary.post(reverse("staffpanel:seminar_clone", kwargs={"pk": source.pk}))

    source.refresh_from_db()
    assert source.talks.count() == 1
    assert source.status == Seminar.Status.PUBLISHED


def test_clone_gets_unique_slug(as_secretary):
    source = make_seminar(timezone.localdate() - timedelta(days=30))

    as_secretary.post(reverse("staffpanel:seminar_clone", kwargs={"pk": source.pk}))

    slugs = list(Seminar.objects.values_list("slug", flat=True))
    assert len(slugs) == len(set(slugs))


# --- Показать/скрыть ----------------------------------------------------------


def test_toggle_hides_published_seminar(as_secretary):
    seminar = make_seminar(timezone.localdate() + timedelta(days=5))

    as_secretary.post(reverse("staffpanel:seminar_toggle", kwargs={"pk": seminar.pk}))

    seminar.refresh_from_db()
    assert seminar.status == Seminar.Status.HIDDEN


def test_toggle_publishes_draft(as_secretary):
    seminar = make_seminar(timezone.localdate() + timedelta(days=5), status=Seminar.Status.DRAFT)

    as_secretary.post(reverse("staffpanel:seminar_toggle", kwargs={"pk": seminar.pk}))

    seminar.refresh_from_db()
    assert seminar.status == Seminar.Status.PUBLISHED


def test_toggle_ignores_foreign_referer(as_secretary):
    """Referer — не место для цели редиректа: иначе это открытый редирект."""
    seminar = make_seminar(timezone.localdate() + timedelta(days=5))

    response = as_secretary.post(
        reverse("staffpanel:seminar_toggle", kwargs={"pk": seminar.pk}),
        HTTP_REFERER="https://evil.example/phish",
    )

    assert response["Location"] == reverse("staffpanel:seminar_list")


# --- Регистрации --------------------------------------------------------------


def test_registrations_filtered_by_seminar(as_secretary):
    first = make_seminar(timezone.localdate() + timedelta(days=5), suffix="a")
    second = make_seminar(timezone.localdate() + timedelta(days=9), suffix="b")
    make_registration(first, full_name="Первый Участник", email="one@example.org")
    make_registration(second, full_name="Второй Участник", email="two@example.org")

    content = as_secretary.get(
        reverse("staffpanel:registrations") + f"?seminar={first.pk}"
    ).content.decode()

    assert "Первый Участник" in content
    assert "Второй Участник" not in content


def test_registrations_search(as_secretary):
    seminar = make_seminar(timezone.localdate() + timedelta(days=5))
    make_registration(seminar, full_name="Петрова Мария", email="petrova@izmiran.ru")
    make_registration(seminar, full_name="Сидоров Пётр", email="sidorov@cosmos.ru")

    content = as_secretary.get(
        reverse("staffpanel:registrations") + f"?seminar={seminar.pk}&q=izmiran"
    ).content.decode()

    assert "Петрова Мария" in content
    assert "Сидоров Пётр" not in content


def test_pass_status_can_be_updated(as_secretary):
    seminar = make_seminar(timezone.localdate() + timedelta(days=5))
    registration = make_registration(seminar)

    as_secretary.post(
        reverse("staffpanel:pass_status", kwargs={"pk": registration.pk}),
        {"pass_status": Registration.PassStatus.ISSUED},
    )

    registration.refresh_from_db()
    assert registration.pass_status == Registration.PassStatus.ISSUED


def test_bogus_pass_status_is_ignored(as_secretary):
    seminar = make_seminar(timezone.localdate() + timedelta(days=5))
    registration = make_registration(seminar, pass_status=Registration.PassStatus.PENDING)

    as_secretary.post(
        reverse("staffpanel:pass_status", kwargs={"pk": registration.pk}),
        {"pass_status": "уже оформлен, честное слово"},
    )

    registration.refresh_from_db()
    assert registration.pass_status == Registration.PassStatus.PENDING


# --- Выгрузка -----------------------------------------------------------------


def test_csv_opens_in_excel_without_mojibake(as_secretary):
    """BOM и точка с запятой — иначе Excel в русской локали ломает файл."""
    seminar = make_seminar(timezone.localdate() + timedelta(days=5))
    make_registration(seminar, full_name="Петрова Мария Сергеевна", organization="ИЗМИРАН")

    response = as_secretary.get(
        reverse("staffpanel:registrations_export", kwargs={"fmt": "csv"}) + f"?seminar={seminar.pk}"
    )

    assert response.status_code == 200
    body = response.content
    assert body.startswith(b"\xef\xbb\xbf"), "нет BOM — Excel прочитает UTF-8 как cp1251"
    text = body.decode("utf-8-sig")
    assert ";" in text.splitlines()[0]
    assert "Петрова Мария Сергеевна" in text
    assert "ИЗМИРАН" in text


def test_csv_respects_active_filters(as_secretary):
    first = make_seminar(timezone.localdate() + timedelta(days=5), suffix="a")
    second = make_seminar(timezone.localdate() + timedelta(days=9), suffix="b")
    make_registration(first, full_name="Нужный Участник", email="one@example.org")
    make_registration(second, full_name="Лишний Участник", email="two@example.org")

    text = as_secretary.get(
        reverse("staffpanel:registrations_export", kwargs={"fmt": "csv"}) + f"?seminar={first.pk}"
    ).content.decode("utf-8-sig")

    assert "Нужный Участник" in text
    assert "Лишний Участник" not in text


def test_xlsx_export_is_a_real_workbook(as_secretary):
    from io import BytesIO

    from openpyxl import load_workbook

    seminar = make_seminar(timezone.localdate() + timedelta(days=5))
    make_registration(seminar, full_name="Петрова Мария Сергеевна")

    response = as_secretary.get(
        reverse("staffpanel:registrations_export", kwargs={"fmt": "xlsx"})
        + f"?seminar={seminar.pk}"
    )

    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.cell(row=1, column=1).value == "ФИО"
    assert sheet.cell(row=2, column=1).value == "Петрова Мария Сергеевна"


def test_export_filename_mentions_seminar_date(as_secretary):
    seminar = make_seminar(timezone.localdate() + timedelta(days=5))
    make_registration(seminar)

    response = as_secretary.get(
        reverse("staffpanel:registrations_export", kwargs={"fmt": "csv"}) + f"?seminar={seminar.pk}"
    )

    assert f"{seminar.date:%Y-%m-%d}" in response["Content-Disposition"]


def test_unknown_export_format_is_404(as_secretary):
    response = as_secretary.get(reverse("staffpanel:registrations_export", kwargs={"fmt": "pdf"}))

    assert response.status_code == 404


def test_export_is_closed_to_outsiders(client, django_user_model):
    django_user_model.objects.create_user(username="guest", password="pw")
    client.force_login(django_user_model.objects.get(username="guest"))

    response = client.get(reverse("staffpanel:registrations_export", kwargs={"fmt": "csv"}))

    assert response.status_code == 403
