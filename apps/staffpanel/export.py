"""Выгрузка регистраций.

Данные персональные, поэтому выгрузка доступна только сотрудникам,
а имя файла содержит дату заседания — чтобы выгруженные списки не
перепутались между собой.
"""

import csv
from io import BytesIO, StringIO

from django.http import HttpResponse
from django.utils.text import slugify
from django.utils.translation import gettext as _

COLUMNS = [
    (_("ФИО"), "full_name"),
    (_("Организация"), "organization"),
    (_("Должность"), "position"),
    (_("Почта"), "email"),
    (_("Формат участия"), "attendance_display"),
    (_("Гражданство"), "citizenship_display"),
    (_("Пропуск"), "pass_display"),
    (_("Заявка подана"), "submitted"),
]


def _rows(registrations):
    for registration in registrations:
        yield {
            "full_name": registration.full_name,
            "organization": registration.organization,
            "position": registration.position or "—",
            "email": registration.email,
            "attendance_display": registration.get_attendance_display(),
            "citizenship_display": registration.citizenship_display,
            "pass_display": registration.get_pass_status_display(),
            "submitted": registration.created_at.strftime("%d.%m.%Y %H:%M"),
        }


def _filename(seminar, extension: str) -> str:
    if seminar is None:
        return f"registrations-all.{extension}"
    # Тема первого доклада: собственного заголовка у заседания нет, а по одной
    # дате в папке выгрузок не разобрать, что за файл.
    parts = ["registrations", f"{seminar.date:%Y-%m-%d}", slugify(seminar.label)[:40]]
    return "-".join(part for part in parts if part) + f".{extension}"


def registrations_csv(registrations, seminar=None) -> HttpResponse:
    """CSV, который открывается в Excel без плясок с кодировкой.

    Два обязательных условия: BOM (иначе Excel на Windows читает UTF-8
    как cp1251 и показывает кракозябры) и точка с запятой в качестве
    разделителя (в русской локали Excel запятая — десятичный знак).
    """
    buffer = StringIO()
    writer = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow([str(label) for label, _key in COLUMNS])
    for row in _rows(registrations):
        writer.writerow([row[key] for _label, key in COLUMNS])

    response = HttpResponse(
        buffer.getvalue().encode("utf-8-sig"), content_type="text/csv; charset=utf-8"
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename(seminar, "csv")}"'
    return response


def registrations_xlsx(registrations, seminar=None) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заявки"

    sheet.append([str(label) for label, _key in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")

    widths = [len(str(label)) for label, _key in COLUMNS]
    for row in _rows(registrations):
        values = [row[key] for _label, key in COLUMNS]
        sheet.append(values)
        widths = [max(w, min(len(str(v)), 60)) for w, v in zip(widths, values, strict=True)]

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width + 3
    sheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename(seminar, "xlsx")}"'
    return response
